import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_simplified_architecture(output_path="ILRDVS_Simplified_Architecture.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    C_BG = RGBColor(248, 250, 252)          # #F8FAFC
    C_NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A
    C_SLATE_BORDER = RGBColor(203, 213, 225)# #CBD5E1
    C_WHITE = RGBColor(255, 255, 255)
    C_DARK_TEXT = RGBColor(15, 23, 42)      # #0F172A
    C_MUTED_TEXT = RGBColor(71, 85, 105)    # #475569

    # Tier Theme Colors
    TIER_COLORS = [
        {"ribbon": RGBColor(30, 58, 138),  "border": RGBColor(59, 130, 246),  "bg": RGBColor(239, 246, 255), "arrow": RGBColor(59, 130, 246)},  # 1. Blue
        {"ribbon": RGBColor(2, 132, 199),  "border": RGBColor(14, 165, 233),  "bg": RGBColor(240, 249, 255), "arrow": RGBColor(14, 165, 233)},  # 2. Sky
        {"ribbon": RGBColor(4, 120, 87),   "border": RGBColor(16, 185, 129),  "bg": RGBColor(236, 253, 245), "arrow": RGBColor(16, 185, 129)},  # 3. Green
        {"ribbon": RGBColor(217, 119, 6),  "border": RGBColor(245, 158, 11),  "bg": RGBColor(255, 251, 235), "arrow": RGBColor(245, 158, 11)},  # 4. Amber
        {"ribbon": RGBColor(124, 58, 237), "border": RGBColor(168, 85, 247),  "bg": RGBColor(250, 245, 255), "arrow": RGBColor(168, 85, 247)},  # 5. Purple
    ]

    slide = prs.slides.add_slide(blank_layout)

    # 1. Slide Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_BG
    bg.line.fill.background()

    # 2. Top Header Bar
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.95))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = C_NAVY_DARK
    top_bar.line.fill.background()

    tf = top_bar.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.6)
    p = tf.paragraphs[0]
    p.text = "ILRDVS — SIMPLIFIED SYSTEM ARCHITECTURE"
    p.font.name = "Trebuchet MS"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p_sub = tf.add_paragraph()
    p_sub.text = "5-Tier High-Level Pipeline: Citizen Ingestion ➔ AI Extraction ➔ Officer Sanction ➔ Secure QR"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(10)
    p_sub.font.color.rgb = RGBColor(147, 197, 253)

    # Top Right Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.2), Inches(0.22), Inches(2.5), Inches(0.48))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(217, 119, 6)
    badge.line.fill.background()
    btf = badge.text_frame
    btf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = btf.paragraphs[0]
    bp.text = "SMART INDIA HACKATHON"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.name = "Arial"
    bp.font.size = Pt(9.5)
    bp.font.bold = True
    bp.font.color.rgb = C_WHITE

    # Stripe
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.95), Inches(13.333), Inches(0.05))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = RGBColor(217, 119, 6)
    stripe.line.fill.background()

    # -------------------------------------------------------------------------
    # 5 SIMPLIFIED TIERS DEFINITION
    # -------------------------------------------------------------------------
    tiers_data = [
        {
            "num": "1",
            "name": "USER / CLIENT",
            "tech": "Next.js 14",
            "cards": [
                ("Citizen Portal (/citizen)", "OTP Login • Deed Upload • Application Tracking"),
                ("Officer Console (/officer)", "Verification Queue • Statutory Inspection"),
                ("Public QR Verify (/verify)", "Instant 3-Tier Document Authenticity Check")
            ]
        },
        {
            "num": "2",
            "name": "API & SECURITY",
            "tech": "Node.js + Express",
            "cards": [
                ("REST APIs & Router", "Modular Endpoints • Centralized Request Gateway"),
                ("RBAC & Multi-Factor Auth", "JWT Cookies • DLT SMS (MSG91) • HMAC Email OTP"),
                ("Zod Guards & Multer Ingest", "Strict Schema Validation • Scanned File Vault")
            ]
        },
        {
            "num": "3",
            "name": "AI PIPELINE",
            "tech": "Python FastAPI",
            "cards": [
                ("OpenCV Vision", "CLAHE 3.5 Contrast Boost • Faint Ink Rescue • Deskew"),
                ("Multilingual OCR", "EasyOCR + Tesseract (Marathi Devanagari + English)"),
                ("Cadastral NER Extractor", "Extracts Owner, Survey/Gat #, Khata Account, Area")
            ]
        },
        {
            "num": "4",
            "name": "REVIEW & STORE",
            "tech": "HITL + MongoDB 8",
            "cards": [
                ("Split-Screen Workstation", "Original Scan vs AI Extracted Entities Review"),
                ("Statutory Officer Sign", "Dual Digital Signatures (DSC-VER & DSC-OFF)"),
                ("MongoDB 8 Persistence", "LandRecords (ULPIN-Indexed) + Immutable AuditLog")
            ]
        },
        {
            "num": "5",
            "name": "SECURE SEAL",
            "tech": "QR + Cryptography",
            "cards": [
                ("Digital Cloud Anchor", "Unique Doc ID (DOC-MH-2026-XXXX) + SHA-256 Hash"),
                ("+ Physical Sticker Seal", "80mm Tamper-Evident Adhesive Sticker + Secret PIN"),
                ("➔ 3-Tier Outcome", "[PASS] Valid Match | [WARN] Digital | [ALERT] Tamper")
            ]
        }
    ]

    start_y = 1.18
    tier_height = 0.88
    gap_y = 0.28
    tier_width = 12.133
    ribbon_width = 2.45

    for i, tdata in enumerate(tiers_data):
        y_pos = start_y + i * (tier_height + gap_y)
        cfg = TIER_COLORS[i]

        # 1. Main Tier Outer Card
        outer_card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.6), Inches(y_pos), Inches(tier_width), Inches(tier_height)
        )
        outer_card.fill.solid()
        outer_card.fill.fore_color.rgb = cfg["bg"]
        outer_card.line.color.rgb = cfg["border"]
        outer_card.line.width = Pt(1.2)

        # 2. Left Ribbon Header
        ribbon = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.6), Inches(y_pos), Inches(ribbon_width), Inches(tier_height)
        )
        ribbon.fill.solid()
        ribbon.fill.fore_color.rgb = cfg["ribbon"]
        ribbon.line.fill.background()

        rtf = ribbon.text_frame
        rtf.vertical_anchor = MSO_ANCHOR.MIDDLE
        rtf.margin_left = Inches(0.18)
        rtf.margin_right = Inches(0.1)

        rp1 = rtf.paragraphs[0]
        rp1.text = f"{tdata['num']}. {tdata['name']}"
        rp1.font.name = "Trebuchet MS"
        rp1.font.bold = True
        rp1.font.size = Pt(11)
        rp1.font.color.rgb = C_WHITE

        rp2 = rtf.add_paragraph()
        rp2.text = tdata["tech"]
        rp2.font.name = "Arial"
        rp2.font.size = Pt(8.5)
        rp2.font.color.rgb = RGBColor(226, 232, 240)

        # 3. Inside 3 Step/Pill Cards
        pill_start_x = 3.25
        pill_gap = 0.18
        pill_width = 3.02
        pill_height = 0.64
        pill_y = y_pos + 0.12

        for c_idx, (card_title, card_sub) in enumerate(tdata["cards"]):
            cx = pill_start_x + c_idx * (pill_width + pill_gap)
            pill = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(cx), Inches(pill_y), Inches(pill_width), Inches(pill_height)
            )
            pill.fill.solid()
            pill.fill.fore_color.rgb = C_WHITE
            pill.line.color.rgb = cfg["border"]
            pill.line.width = Pt(0.9)

            ptf = pill.text_frame
            ptf.word_wrap = True
            ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
            ptf.margin_left = Inches(0.12)
            ptf.margin_right = Inches(0.12)
            ptf.margin_top = Inches(0.04)
            ptf.margin_bottom = Inches(0.04)

            cp1 = ptf.paragraphs[0]
            cp1.text = card_title
            cp1.font.name = "Trebuchet MS"
            cp1.font.bold = True
            cp1.font.size = Pt(9.5)
            cp1.font.color.rgb = cfg["ribbon"]

            cp2 = ptf.add_paragraph()
            cp2.text = card_sub
            cp2.font.name = "Arial"
            cp2.font.size = Pt(7.8)
            cp2.font.color.rgb = C_MUTED_TEXT

        # 4. Downward Connecting Arrow (if not last tier)
        if i < len(tiers_data) - 1:
            arr_y = y_pos + tier_height + 0.04
            arr = slide.shapes.add_shape(
                MSO_SHAPE.DOWN_ARROW,
                Inches(6.6), Inches(arr_y), Inches(0.24), Inches(0.20)
            )
            arr.fill.solid()
            arr.fill.fore_color.rgb = cfg["arrow"]
            arr.line.fill.background()

    # Footer
    f_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(7.05), Inches(12.133), Inches(0.02))
    f_line.fill.solid()
    f_line.fill.fore_color.rgb = C_SLATE_BORDER
    f_line.line.fill.background()

    txBox = slide.shapes.add_textbox(Inches(0.6), Inches(7.08), Inches(8.5), Inches(0.35))
    tf_f = txBox.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "ILRDVS • Simplified System Architecture • Clear 5-Tier Governance Flow"
    p_f.font.name = "Arial"
    p_f.font.size = Pt(9)
    p_f.font.color.rgb = C_MUTED_TEXT

    txBox_r = slide.shapes.add_textbox(Inches(9.5), Inches(7.08), Inches(3.233), Inches(0.35))
    tf_r = txBox_r.text_frame
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.text = "100% Native PowerPoint Shapes (Fully Editable)"
    p_r.font.name = "Arial"
    p_r.font.size = Pt(9)
    p_r.font.bold = True
    p_r.font.color.rgb = RGBColor(30, 58, 138)

    # Save
    prs.save(output_path)
    print(f"Simplified architecture presentation created successfully: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_simplified_architecture()
